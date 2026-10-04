"""Report export (Excel + CSV). Writes only to outputs/ (or a given path)."""
from __future__ import annotations

import io
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config as C
from .proposals import accepted_corrections

CASE_COLUMNS = [
    "case_id", "person_id", "nom", "case_type", "field", "assignment", "src_row", "dst_row",
    "match_status", "match_confidence", "source_fields", "source_value", "destination_value",
    "source_normalized", "destination_normalized", "expected_value", "alt_expected_value", "stage1_result",
    "rule_id", "rule_name", "rule_type", "rule_outcome", "rule_note", "certainty", "deterministic_verdict",
    "verdict", "confidence", "priority", "priority_score", "explanation", "ai_hypothesis", "ai_suggested_action",
    "ai_suggested_verdict", "proposed_action", "pattern_id", "systemic", "ambiguity_id", "quality_flag", "inputs", "context",
    "evidence", "mapping_ref", "confidence_detail", "priority_detail", "human_decision", "human_decision_source",
    "human_comment", "human_decided_at", "reviewed_verdict",
]
AUDIT_COLUMNS = ["case_id", "person_id", "field", "case_type", "source_value", "destination_value", "expected_value",
                 "rule_id", "rule_type", "deterministic_verdict", "verdict", "pattern_id", "human_decision",
                 "human_decision_source", "human_comment", "human_decided_at", "reviewed_verdict", "evidence"]


def audit_log(result) -> pd.DataFrame:
    """Engine verdict next to the human decision for every non-conforming or reviewed case."""
    c = result.cases
    mask = c["verdict"].isin([C.ANOMALIE, C.A_INVESTIGUER]) | (c["human_decision"] != "")
    cols = [x for x in AUDIT_COLUMNS if x in c.columns]
    log = c.loc[mask, cols].copy()
    if not result.proposals.empty:
        status = result.proposals.groupby("case_id")["status"].agg(lambda s: ", ".join(sorted(set(s))))
        log["proposal_status"] = log["case_id"].map(status).fillna("")
    return log
VERDICT_FILL = {C.CONFORME: "E8F5E9", C.ECART_JUSTIFIE: "E3F2FD", C.ANOMALIE: "FFEBEE", C.A_INVESTIGUER: "FFF8E1"}
MAX_XLSX_CASES = 100_000


def _cases_view(cases: pd.DataFrame) -> pd.DataFrame:
    cols = [c for c in CASE_COLUMNS if c in cases.columns]
    return cases[cols].sort_values(["priority_score", "person_id"], ascending=[False, True])


def _relative(path) -> str:
    """Repository-relative path for public reports (no machine-specific absolute paths)."""
    try:
        return Path(path).resolve().relative_to(C.PROJECT_ROOT.resolve()).as_posix()
    except ValueError:
        return Path(path).name


def summary_frame(result) -> pd.DataFrame:
    s = result.summary
    rows = [
        ("Jeu de données", _relative(s["dataset"])),
        ("Date d'exécution", datetime.now().isoformat(timespec="seconds")),
        ("Méthodologie", "Niveau 1 comparaison brute/normalisation → Niveau 2 règles déterministes (Mapping.xlsx) → "
                         "Niveau 3 analyse assistée des cas ambigus"),
        ("Personnes", s["n_persons"]),
        ("Affectations source", s["n_source_assignments"]),
        ("Lignes destination", s["n_destination_rows"]),
        ("Champs corroborés (Mapping.xlsx)", s["n_fields_corroborated"]),
        ("Cas analysés", s["n_cases"]),
        ("  dont comparaisons de champs", s["n_field_comparisons"]),
        ("CONFORME", s["n_CONFORME"]),
        ("ECART_JUSTIFIE", s["n_ECART_JUSTIFIE"]),
        ("ANOMALIE", s["n_ANOMALIE"]),
        ("A_INVESTIGUER", s["n_A_INVESTIGUER"]),
        ("Cas regroupés en motifs systémiques", s["n_systemic_cases"]),
        ("Groupes systémiques / ambiguïtés de règle (1 décision chacun)", s["n_systemic_groups"]),
        ("Cas individuels nécessitant une attention humaine", s["n_attention_individual"]),
        ("TOTAL éléments nécessitant une attention humaine", s["n_human_attention"]),
        ("Appariements : appariés / ambigus / source seule / destination seule",
         f"{s['n_matched']} / {s['n_ambiguous_match']} / {s['n_source_only']} / {s['n_dest_only']}"),
        ("Verdicts par niveau : comparaison brute / règle déterministe / assisté IA",
         f"{s['n_rule_type_' + C.RAW]} / {s['n_rule_type_' + C.DETERMINISTIC]} / {s['n_rule_type_' + C.AI_ASSISTED]}"),
        ("Propositions de correction", s["n_proposals"]),
        ("Cas revus par un expert", s["n_reviewed"]),
    ]
    return pd.DataFrame(rows, columns=["Indicateur", "Valeur"])


def field_pivot(cases: pd.DataFrame) -> pd.DataFrame:
    p = pd.crosstab(cases["field"], cases["verdict"]).reindex(columns=C.VERDICTS, fill_value=0)
    p["TOTAL"] = p.sum(axis=1)
    return p.reset_index()


def report_sheets(result) -> dict:
    cases = result.cases
    view = _cases_view(cases)
    note = None
    if len(view) > MAX_XLSX_CASES:
        view = view[view["verdict"] != C.CONFORME]
        note = "Cas CONFORME omis de la feuille Cases (volume) : voir cases.csv"
    attention = view[view["verdict"].isin([C.ANOMALIE, C.A_INVESTIGUER])]
    summary = summary_frame(result)
    if note:
        summary.loc[len(summary)] = ["Note", note]
    sheets = {
        "Summary": summary,
        "By_Field": field_pivot(cases),
        "Anomalies": attention[attention["verdict"] == C.ANOMALIE],
        "To_Investigate": attention[(attention["verdict"] == C.A_INVESTIGUER) & ~attention["systemic"].astype(bool)],
        "Patterns": result.patterns,
        "Systemic_Cases": attention[attention["systemic"].astype(bool)],
        "Justified": view[view["verdict"] == C.ECART_JUSTIFIE],
        "Proposed_Actions": result.proposals,
        "Accepted_Actions": accepted_corrections(result.proposals),
        "Rule_Ambiguities": result.ambiguities,
        "Audit_Log": audit_log(result),
        "Matching": result.matches,
        "Mapping": result.mapping,
        "Cases": view,
    }
    return sheets


def _format(ws, df: pd.DataFrame):
    ws.freeze_panes = "A2"
    if df.shape[1]:
        ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="37474F")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for i, col in enumerate(df.columns, start=1):
        sample = df[col].astype(str).head(200)
        width = min(max([len(str(col))] + [len(x) for x in sample]) + 2, 60)
        ws.column_dimensions[get_column_letter(i)].width = max(width, 8)
    if "verdict" in df.columns and len(df) <= 20000:
        vcol = list(df.columns).index("verdict") + 1
        for row in ws.iter_rows(min_row=2, min_col=vcol, max_col=vcol):
            for cell in row:
                color = VERDICT_FILL.get(cell.value)
                if color:
                    cell.fill = PatternFill("solid", fgColor=color)


def write_report(result, target) -> None:
    sheets = report_sheets(result)
    with pd.ExcelWriter(target, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df = df.copy()
            for c in df.columns:
                if df[c].dtype == object:
                    df[c] = df[c].map(lambda v: v if not isinstance(v, (list, dict, set)) else str(v))
            df.to_excel(xw, sheet_name=name[:31], index=False)
            _format(xw.sheets[name[:31]], df)


def export_report(result, path=None) -> Path:
    path = Path(path) if path else C.OUTPUT_DIR / "corroboration_report.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    write_report(result, path)
    return path


def report_bytes(result) -> bytes:
    buf = io.BytesIO()
    write_report(result, buf)
    return buf.getvalue()


def export_csv(result, out_dir=None) -> list[Path]:
    out = Path(out_dir) if out_dir else C.OUTPUT_DIR
    out.mkdir(parents=True, exist_ok=True)
    files = {
        "cases.csv": _cases_view(result.cases),
        "anomalies.csv": _cases_view(result.cases[result.cases["verdict"] == C.ANOMALIE]),
        "to_investigate.csv": _cases_view(result.cases[result.cases["verdict"] == C.A_INVESTIGUER]),
        "patterns.csv": result.patterns,
        "proposed_actions.csv": result.proposals,
        "rule_ambiguities.csv": result.ambiguities,
    }
    paths = []
    for name, df in files.items():
        p = out / name
        df.to_csv(p, index=False, encoding="utf-8-sig")
        paths.append(p)
    return paths


def corrected_destination(result) -> tuple[pd.DataFrame, pd.DataFrame]:
    """SIMULATION: copy of the destination extract with every accepted/modified correction applied
    (field updates + candidate rows). The official files and the loaded dataset are never modified.
    Returns (candidate destination, change log with provenance)."""
    dest = result.dataset.destination.copy().astype(object)
    dest.insert(0, "origine_ligne", dest["_excel_row"].map(lambda r: f"ligne {r} du fichier officiel"))
    acc = accepted_corrections(result.proposals)
    changes = []
    for _, r in acc[acc["action_type"] == "CORRIGER_CHAMP"].iterrows():
        idx = dest.index[dest["_excel_row"] == r["dest_row"]]
        if not len(idx) or r["field"] not in dest.columns:
            continue
        before = dest.at[idx[0], r["field"]]
        dest.at[idx[0], r["field"]] = r["final_value"]
        dest.at[idx[0], "origine_ligne"] = f"ligne {r['dest_row']} du fichier officiel (corrigée)"
        changes.append({"type": "MODIFICATION", "person_id": r["person_id"], "ligne_destination": r["dest_row"],
                        "champ": r["field"], "avant": "" if pd.isna(before) else str(before),
                        "apres": r["final_value"], "statut": r["status"], "case_id": r["case_id"],
                        "provenance": r["derivation"]})
    new_rows = []
    for pid, grp in acc[acc["action_type"] == "CREER_AFFECTATION"].groupby("proposal_id", sort=False):
        row = {c: "" for c in dest.columns}
        row.update({f: v for f, v in zip(grp["field"], grp["final_value"]) if f in dest.columns})
        row["origine_ligne"] = f"NOUVELLE ligne candidate ({pid})"
        row["_excel_row"] = None
        new_rows.append(row)
        changes.append({"type": "CREATION", "person_id": grp["person_id"].iloc[0], "ligne_destination": "nouvelle",
                        "champ": f"{len(grp)} champs", "avant": "", "apres": "nouvelle affectation candidate",
                        "statut": ", ".join(sorted(set(grp["status"]))), "case_id": grp["case_id"].iloc[0],
                        "provenance": "proposition de création dérivée du mapping"})
    if new_rows:
        dest = pd.concat([dest, pd.DataFrame(new_rows, columns=dest.columns)], ignore_index=True)
    return dest.drop(columns=["_excel_row"]), pd.DataFrame(
        changes, columns=["type", "person_id", "ligne_destination", "champ", "avant", "apres", "statut", "case_id",
                          "provenance"])


def export_corrected_destination(result, path=None) -> Path:
    """Write the simulated corrected destination to a SEPARATE file (never inside data/)."""
    path = Path(path) if path else C.OUTPUT_DIR / "destination_corrigee_candidate.xlsx"
    if path.resolve().is_relative_to(C.DATA_DIR.resolve()):
        raise ValueError("Refus d'écrire dans data/ : les fichiers officiels sont en lecture seule.")
    dest, changes = corrected_destination(result)
    path.parent.mkdir(parents=True, exist_ok=True)
    readme = pd.DataFrame({"Avertissement": [
        "SIMULATION — copie candidate de l'extraction destination avec les corrections acceptées.",
        "Aucun fichier officiel n'a été modifié.",
        "Chaque valeur modifiée est tracée dans la feuille Modifications (provenance, décision, cas)."]})
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        readme.to_excel(xw, sheet_name="Lisez-moi", index=False)
        changes.to_excel(xw, sheet_name="Modifications", index=False)
        dest.to_excel(xw, sheet_name="Destination_candidate", index=False)
        for name, df in (("Lisez-moi", readme), ("Modifications", changes), ("Destination_candidate", dest)):
            _format(xw.sheets[name], df)
    return path


def export_accepted(result, out_dir=None) -> Path:
    out = Path(out_dir) if out_dir else C.OUTPUT_DIR
    out.mkdir(parents=True, exist_ok=True)
    acc = accepted_corrections(result.proposals)
    p = out / "accepted_corrections.csv"
    acc.to_csv(p, index=False, encoding="utf-8-sig")
    with pd.ExcelWriter(out / "accepted_corrections.xlsx", engine="openpyxl") as xw:
        acc.to_excel(xw, sheet_name="Accepted_Actions", index=False)
        _format(xw.sheets["Accepted_Actions"], acc)
    return p
